"""Read a source CV and render simple, text-based application documents."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from xml.sax.saxutils import escape
from zipfile import ZipFile


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
MAX_CV_CHARACTERS = 100_000
MAX_PDF_PAGES = 50
_INSTALL = "Install application dependencies with: pip install -r requirements-openai.txt"
_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _validate_text(text: str) -> str:
    text = text.strip()
    if not text:
        raise ValueError(
            "The CV contains no extractable text. For a scanned PDF, run OCR first "
            "or provide a text-based PDF, DOCX, Markdown, or UTF-8 text file."
        )
    if len(text) > MAX_CV_CHARACTERS:
        raise ValueError(f"CV text exceeds the {MAX_CV_CHARACTERS:,}-character limit.")
    if _INVALID_XML.search(text):
        raise ValueError("CV text contains unsupported control characters; export a clean text copy.")
    return text


def _docx_module():
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError(f"DOCX support requires python-docx. {_INSTALL}") from exc
    return docx


def _word_blocks(element, parent) -> list[str]:
    """Extract a Word body, header, or footer in paragraph/table order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    blocks = []
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            blocks.append(Paragraph(child, parent).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, parent).rows:
                # A merged cell can appear more than once within a row.
                seen_cells = set()
                cells = []
                for cell in row.cells:
                    if cell._tc not in seen_cells:
                        cells.append(cell.text)
                        seen_cells.add(cell._tc)
                blocks.append("\t".join(cells))
    return blocks


def read_cv(path: Path | str) -> str:
    """Extract bounded plain text from .txt, .md, .pdf, or .docx files.

    Image-only PDFs need OCR before use. Dependencies are imported only for the
    requested format, so a text source CV can be read without document libraries.
    """
    source = Path(path).expanduser()
    suffix = source.suffix.lower()
    if suffix not in {".txt", ".md", ".pdf", ".docx"}:
        raise ValueError("Unsupported CV format. Use .txt, .md, .pdf, or .docx.")
    if source.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("CV file exceeds the 10 MiB size limit.")
    payload = source.read_bytes()
    if len(payload) > MAX_FILE_BYTES:
        raise ValueError("CV file exceeds the 10 MiB size limit.")

    if suffix in {".txt", ".md"}:
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Save the text CV using UTF-8 encoding and try again.") from exc
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError(f"PDF reading requires pypdf. {_INSTALL}") from exc
        try:
            reader = PdfReader(BytesIO(payload))
            if reader.is_encrypted:
                raise ValueError("The CV PDF is encrypted. Export an unencrypted copy first.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError(f"CV PDF exceeds the {MAX_PDF_PAGES}-page limit.")
            pages = []
            character_count = 0
            for page in reader.pages:
                page_text = page.extract_text() or ""
                character_count += len(page_text)
                if character_count > MAX_CV_CHARACTERS:
                    raise ValueError(f"CV text exceeds the {MAX_CV_CHARACTERS:,}-character limit.")
                pages.append(page_text)
            text = "\n\n".join(pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read the PDF CV. Export a fresh text-based PDF and retry.") from exc
    else:
        docx = _docx_module()
        try:
            with ZipFile(BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_DOCUMENT_BYTES:
                    raise ValueError("Expanded DOCX content exceeds the 50 MiB size limit.")
            document = docx.Document(BytesIO(payload))
            headers, footers = [], []
            seen_regions = set()
            for section in document.sections:
                variants = [""]
                if section.different_first_page_header_footer:
                    variants.append("first_page_")
                if document.settings.odd_and_even_pages_header_footer:
                    variants.append("even_page_")
                for kind, collected in (("header", headers), ("footer", footers)):
                    for variant in variants:
                        region = getattr(section, variant + kind)
                        # Linked sections share a region already visited. Skipping
                        # them also avoids creating an absent header during a read.
                        if region.is_linked_to_previous:
                            continue
                        content = "\n".join(_word_blocks(region._element, region)).strip()
                        if content and content not in seen_regions:
                            collected.append(content)
                            seen_regions.add(content)
            body = _word_blocks(document.element.body, document)
            text = "\n".join([*headers, *body, *footers])
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Could not read the DOCX CV. Export a fresh DOCX file and retry.") from exc
    return _validate_text(text)


def _blocks(markdown: str) -> list[tuple[str, int, str]]:
    """Recognize only headings and list prefixes; all inline content stays text."""
    result = []
    paragraph = []

    def flush() -> None:
        if paragraph:
            result.append(("paragraph", 0, " ".join(paragraph)))
            paragraph.clear()

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        bullet = re.match(r"^(?:[-+*]|\d+[.)])\s+(.+)$", line)
        if not line:
            flush()
        elif heading:
            flush()
            result.append(("heading", len(heading[1]), heading[2]))
        elif bullet:
            flush()
            result.append(("bullet", 0, bullet[1]))
        else:
            paragraph.append(line)
    flush()
    return result


def _pdf_fonts():
    """Prefer installed Unicode fonts, with ReportLab's bundled Vera as fallback."""
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    reportlab_fonts = Path(reportlab.__file__).resolve().parent / "fonts"
    candidates = [
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ),
        (
            Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
            Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
        ),
        (reportlab_fonts / "Vera.ttf", reportlab_fonts / "VeraBd.ttf"),
    ]
    for regular, bold in candidates:
        if regular.is_file() and bold.is_file():
            regular_font = TTFont("ApplicationCV", str(regular))
            bold_font = TTFont("ApplicationCV-Bold", str(bold))
            pdfmetrics.registerFont(regular_font)
            pdfmetrics.registerFont(bold_font)
            return regular_font, bold_font
    raise RuntimeError("PDF export needs a TrueType font. Install the DejaVu Sans font family.")


def export_cv(markdown: str, output_dir: Path | str) -> dict[str, str]:
    """Write cv.md, cv.docx and cv.pdf; return absolute paths for all three.

    Uses one column, ordinary body text, headings and bullets for ATS readability.
    HTML and links are always literal text: no resource is fetched or executed.
    All formats are rendered successfully before destination files are written.
    """
    markdown = _validate_text(markdown)
    blocks = _blocks(markdown)
    docx = _docx_module()
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate
    except ImportError as exc:
        raise RuntimeError(f"PDF export requires reportlab. {_INSTALL}") from exc

    from docx.shared import Inches, Pt, RGBColor

    regular_font, bold_font = _pdf_fonts()
    # An unsupported character should never silently become a blank/tofu glyph.
    for kind, _, value in blocks:
        font = bold_font if kind == "heading" else regular_font
        unsupported = sorted({c for c in value if not c.isspace() and ord(c) not in font.face.charToGlyph})
        if unsupported:
            labels = ", ".join(f"U+{ord(char):04X}" for char in unsupported[:8])
            raise ValueError(
                f"The PDF font cannot render these CV characters: {labels}. "
                "Use supported characters or install DejaVu Sans and retry."
            )

    document = docx.Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.65)
        section.page_width = Inches(A4[0] / inch)
        section.page_height = Inches(A4[1] / inch)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Title"].font.size = Pt(18)
    for level, size in ((1, 13), (2, 11), (3, 10)):
        document.styles[f"Heading {level}"].font.size = Pt(size)

    paragraph_style = ParagraphStyle(
        "CVBody", fontName=regular_font.fontName, fontSize=10, leading=14,
        textColor=colors.black, alignment=TA_LEFT, spaceAfter=5,
    )
    bullet_style = ParagraphStyle(
        "CVBullet", parent=paragraph_style, leftIndent=12, firstLineIndent=0,
        bulletIndent=0, bulletFontName=regular_font.fontName, bulletFontSize=10,
    )
    story = []
    for kind, level, value in blocks:
        if kind == "heading":
            size = 18 if level == 1 else 13 if level == 2 else 11
            style = ParagraphStyle(
                f"CVHeading{level}", parent=paragraph_style,
                fontName=bold_font.fontName, fontSize=size, leading=size + 4,
                spaceBefore=8, spaceAfter=6, keepWithNext=True,
            )
            docx_style = "Title" if level == 1 else f"Heading {min(level - 1, 3)}"
        elif kind == "bullet":
            style = bullet_style
            docx_style = "List Bullet"
        else:
            style = paragraph_style
            docx_style = "Normal"
        document.add_paragraph(value, style=docx_style)
        story.append(Paragraph(escape(value), style, bulletText="•" if kind == "bullet" else None))

    docx_buffer = BytesIO()
    document.save(docx_buffer)
    pdf_buffer = BytesIO()
    pdf = SimpleDocTemplate(
        pdf_buffer, pagesize=A4, leftMargin=0.65 * inch, rightMargin=0.65 * inch,
        topMargin=0.65 * inch, bottomMargin=0.65 * inch, title="Curriculum Vitae",
        author="",
    )
    pdf.build(story)

    destination = Path(output_dir).expanduser().resolve()
    destination.mkdir(parents=True, exist_ok=True)
    rendered = {
        "markdown": ("cv.md", (markdown + "\n").encode("utf-8")),
        "docx": ("cv.docx", docx_buffer.getvalue()),
        "pdf": ("cv.pdf", pdf_buffer.getvalue()),
    }
    paths = {}
    for key, (name, content) in rendered.items():
        path = destination / name
        path.write_bytes(content)
        paths[key] = str(path)
    return paths
