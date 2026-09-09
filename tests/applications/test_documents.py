import builtins
from pathlib import Path

import pytest

from app.applications.documents import MAX_CV_CHARACTERS, export_cv, read_cv


CV = """# Jérôme Müller

jerome@example.test | London | https://example.test/profile

## Professional experience

### Data Engineer — Example Ltd

- Built Python pipelines and saved £2,000 & €500 annually.
- Literal <b>XML</b> and <img src=\"https://example.test/a.png\"/> remain text.

## Education

MSc Computer Science, 2020
"""


def test_export_round_trip_keeps_text_unicode_and_literal_markup(tmp_path):
    from docx import Document

    paths = export_cv(CV, tmp_path / "application")

    assert set(paths) == {"markdown", "docx", "pdf"}
    assert all(Path(path).is_absolute() and Path(path).is_file() for path in paths.values())
    assert read_cv(paths["markdown"]) == CV.strip()
    for extension in ("docx", "pdf"):
        # Text must be selectable/extractable in the documents an employer gets.
        text = " ".join(read_cv(paths[extension]).split())
        for expected in (
            "Jérôme Müller", "jerome@example.test", "Professional experience",
            "Data Engineer — Example Ltd", "£2,000 & €500", "<b>XML</b>",
            '<img src="https://example.test/a.png"/>', "MSc Computer Science, 2020",
        ):
            assert expected in text
        assert "##" not in text

    document = Document(paths["docx"])
    assert not document.tables
    assert document.paragraphs[0].style.name == "Title"
    assert any(p.style.name == "Heading 1" for p in document.paragraphs)
    assert sum(p.style.name == "List Bullet" for p in document.paragraphs) == 2


def test_read_docx_preserves_paragraph_and_table_order(tmp_path):
    from docx import Document

    path = tmp_path / "table_cv.docx"
    document = Document()
    document.add_paragraph("Candidate Name")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python and SQL"
    table.cell(0, 1).text = "Example employer"
    document.add_paragraph("Education")
    document.save(path)

    assert read_cv(path) == "Candidate Name\nPython and SQL\tExample employer\nEducation"


def test_read_docx_includes_contact_headers_and_footers_once(tmp_path):
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Inches

    path = tmp_path / "header_cv.docx"
    document = Document()
    first = document.sections[0]
    first.header.paragraphs[0].text = "Candidate Name | candidate@example.test"
    table = first.footer.add_table(rows=1, cols=1, width=Inches(6))
    table.cell(0, 0).text = "https://example.test/portfolio"
    document.add_paragraph("Experience")
    linked_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    assert linked_section.header.is_linked_to_previous
    document.add_paragraph("Education")
    document.save(path)

    text = read_cv(path)

    assert text.count("candidate@example.test") == 1
    assert text.count("https://example.test/portfolio") == 1
    assert text.index("Candidate Name") < text.index("Experience") < text.index("Education")
    assert text.index("Education") < text.index("https://example.test/portfolio")


def test_empty_pdf_requires_text_or_ocr(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(path)

    with pytest.raises(ValueError, match="OCR"):
        read_cv(path)


def test_encrypted_pdf_requests_unencrypted_copy(tmp_path):
    from pypdf import PdfWriter

    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.encrypt("secret")
    writer.write(path)

    with pytest.raises(ValueError, match="unencrypted"):
        read_cv(path)


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_read_utf8_bom_and_reject_empty_or_oversized_text(tmp_path, suffix):
    path = tmp_path / f"cv{suffix}"
    path.write_text("\ufeff  Candidate résumé\n", encoding="utf-8")
    assert read_cv(path) == "Candidate résumé"
    path.write_text(" \n\t", encoding="utf-8")
    with pytest.raises(ValueError, match="no extractable text"):
        read_cv(path)
    path.write_text("a" * (MAX_CV_CHARACTERS + 1), encoding="utf-8")
    with pytest.raises(ValueError, match="character limit"):
        read_cv(path)


def test_rejects_unsupported_format_and_bad_encoding(tmp_path):
    with pytest.raises(ValueError, match="Unsupported CV format"):
        read_cv(tmp_path / "cv.html")
    path = tmp_path / "cv.txt"
    path.write_bytes(b"\xff\xfeCandidate")
    with pytest.raises(ValueError, match="UTF-8"):
        read_cv(path)


@pytest.mark.parametrize("module_name", ["docx", "reportlab"])
def test_export_missing_dependency_is_actionable_and_writes_no_files(tmp_path, monkeypatch, module_name):
    original_import = builtins.__import__

    def missing_import(name, *args, **kwargs):
        if name == module_name or name.startswith(module_name + "."):
            raise ImportError(f"No module named {module_name}")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", missing_import)
    with pytest.raises(RuntimeError, match="pip install -r requirements-openai.txt"):
        export_cv(CV, tmp_path / "output")
    assert not (tmp_path / "output").exists()


def test_export_rejects_control_characters_before_writing(tmp_path):
    with pytest.raises(ValueError, match="control characters"):
        export_cv("Candidate\x00Name", tmp_path / "output")
    assert not (tmp_path / "output").exists()
